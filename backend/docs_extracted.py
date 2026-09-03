import io
import re
import os
import PyPDF2
from pdfminer.high_level import extract_text as pdfminer_extract_text


class DOCSExtractor:
    def __init__(self, file_buffer, filename=None):
        self.file_buffer = file_buffer
        self.filename = filename
        self.text = None
    
    def _is_docx(self):
        """Check if the file is a DOCX based on filename or magic bytes"""
        if self.filename:
            return self.filename.lower().endswith('.docx')
        if len(self.file_buffer) > 4:
            return self.file_buffer[:2] == b'PK'
        return False
    
    def extract_text(self):
        """Extract text from PDF or DOCX buffer"""
        try:
            if self._is_docx():
                return self._extract_from_docx()
            else:
                return self._extract_from_pdf()
        except Exception as e:
            print(f"Error extracting text: {e}")
            return None
    
    def _extract_from_pdf(self):
        """Extract text from PDF buffer"""
        try:
            try:
                pdf_file = io.BytesIO(self.file_buffer)
                text = pdfminer_extract_text(pdf_file)
                if text and len(text) > 100:
                    self.text = text
                    return text
            except Exception as e:
                print(f"pdfminer extraction failed: {e}")
            
            pdf_file = io.BytesIO(self.file_buffer)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            text = ""
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            self.text = text
            return text
        except Exception as e:
            print(f"Error extracting text from PDF: {e}")
            return None
    
    def _extract_from_docx(self):
        """Extract text from DOCX buffer"""
        try:
            try:
                from docx import Document
            except ImportError:
                print("python-docx not installed. Installing...")
                import subprocess
                import sys
                subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
                from docx import Document
            
            import tempfile
            with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp_file:
                tmp_file.write(self.file_buffer)
                tmp_path = tmp_file.name
            
            doc = Document(tmp_path)
            
            text = ""
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text += paragraph.text + "\n"
            
            # Extract text from tables with better formatting
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text:
                            # Clean up cell text - remove extra newlines within cell
                            cell_text = ' '.join(cell.text.split())
                            row_text.append(cell_text)
                    if row_text:
                        text += " | ".join(row_text) + "\n"
            
            try:
                os.unlink(tmp_path)
            except:
                pass
            
            self.text = text
            return text
            
        except Exception as e:
            print(f"Error extracting text from DOCX: {e}")
            return None
    
    def extract_title(self):
        """Extract main title - find it in the table structure"""
        if not self.text:
            return None
        
        patterns = [
            r'1\.\s*Title\s+of\s+the\s+Extension\s+Project\s+Paper\s*[|:]\s*([^\n|]+)',
            r'1\.\s*Title\s+of\s+the\s+Extension\s+Project\s+Paper\s*\n\s*([^\n]+)',
            r'Title\s+of\s+the\s+Extension\s+Project\s+Paper\s*\|\s*([^\n]+)',
        ]
        
        for pattern in patterns:
            match = re.search(pattern, self.text, re.IGNORECASE)
            if match:
                title = match.group(1).strip()
                title = re.sub(r'\s*(?:2\.|Author|Note:|Name and Email|Paper Category|Thematic Area).*$', '', title, flags=re.IGNORECASE)
                if title and len(title) > 5 and not title.startswith('['):
                    return title
        
        lines = self.text.split('\n')
        title_lines = []
        
        for i, line in enumerate(lines):
            line_stripped = line.strip()
            
            if re.search(r'1\.\s*Title\s+of\s+the\s+Extension\s+Project\s+Paper', line_stripped, re.IGNORECASE):
                for j in range(i + 1, min(i + 10, len(lines))):
                    next_line = lines[j].strip()
                    if not next_line:
                        continue
                    if re.search(r'2\.|Author|Note:|Use an asterisk|paper presenter', next_line, re.IGNORECASE):
                        break
                    if re.search(r'PhD|University|State University|College', next_line, re.IGNORECASE):
                        break
                    if len(next_line) > 10 and not next_line.startswith('['):
                        title_lines.append(next_line)
                break
        
        if title_lines:
            title_text = ' '.join(title_lines)
            title_text = re.sub(r'\s*(?:2\.|Author|Note:|Name and Email|Paper Category|Thematic Area).*$', '', title_text, flags=re.IGNORECASE)
            return title_text
        
        return None
    
    def extract_title_english(self):
        """Extract English version of title"""
        patterns = [
            r'English\s+Version[:\s]*([^\n]+)',
            r'English\s+Title[:\s]*([^\n]+)',
        ]
        for pattern in patterns:
            match = re.search(pattern, self.text, re.IGNORECASE)
            if match:
                return match.group(1).strip()
        return None
    
    def extract_authors(self):
        """Extract authors - get the full author names and affiliations"""
        if not self.text:
            return None
        
        # Method 1: Look for table format with | separator
        pattern = r'2\.\s*Author/s\s+and\s+Institutional\s+Affiliation/s\s*\|\s*([^\n]+)'
        match = re.search(pattern, self.text, re.IGNORECASE)
        if match:
            authors_text = match.group(1).strip()
            # Clean up
            authors_text = re.sub(r'\s*(?:3\.|Name and Email|Paper Category|Thematic Area).*$', '', authors_text, flags=re.IGNORECASE)
            authors_text = ' '.join(authors_text.split())
            if authors_text and len(authors_text) > 5:
                return self._parse_authors(authors_text)
        
        # Method 2: Look for the pattern with "/" separator
        author_pattern = r'([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*\.?\s*[A-Z]\.?\s*[A-Z][a-z]+(?:,\s*(?:PhD|Dr\.)\s*\*?)\s*(?:/\s*[^\n]+)?)'
        match = re.search(author_pattern, self.text, re.IGNORECASE | re.DOTALL)
        if match:
            authors_text = match.group(1).strip()
            authors_text = ' '.join(authors_text.split())
            return self._parse_authors(authors_text)
        
        # Method 3: Look for the authors section by finding the label
        lines = self.text.split('\n')
        for i, line in enumerate(lines):
            if re.search(r'2\.\s*Author/s\s+and\s+Institutional\s+Affiliation/s', line, re.IGNORECASE):
                if '|' in line:
                    parts = line.split('|', 1)
                    if len(parts) == 2:
                        authors_text = parts[1].strip()
                        authors_text = re.sub(r'\s*(?:3\.|Name and Email|Paper Category|Thematic Area).*$', '', authors_text, flags=re.IGNORECASE)
                        authors_text = ' '.join(authors_text.split())
                        if authors_text and len(authors_text) > 5:
                            return self._parse_authors(authors_text)
                else:
                    author_parts = []
                    for j in range(i + 1, min(i + 10, len(lines))):
                        next_line = lines[j].strip()
                        if not next_line:
                            continue
                        if re.search(r'3\.|Name and Email|Paper Category|Thematic Area', next_line, re.IGNORECASE):
                            break
                        if '|' in next_line:
                            parts = next_line.split('|', 1)
                            if len(parts) == 2:
                                author_parts.append(parts[1].strip())
                            else:
                                author_parts.append(next_line)
                        else:
                            author_parts.append(next_line)
                    
                    if author_parts:
                        authors_text = ' '.join(author_parts)
                        authors_text = re.sub(r'\s*(?:3\.|Name and Email|Paper Category|Thematic Area).*$', '', authors_text, flags=re.IGNORECASE)
                        authors_text = ' '.join(authors_text.split())
                        if authors_text and len(authors_text) > 5:
                            return self._parse_authors(authors_text)
                break
        
        return None
    
    def _parse_authors(self, authors_text):
        """Helper method to parse authors text - clean up asterisk but keep PhD"""
        if not authors_text:
            return None
        
        # Store the original text for full_text
        full_text = authors_text
        
        # Extract name and affiliation
        name_part = authors_text
        affiliation_part = ''
        
        # Check if there's a "/" separator
        if '/' in authors_text:
            parts = authors_text.split('/', 1)
            name_part = parts[0].strip()
            affiliation_part = parts[1].strip() if len(parts) > 1 else ''
        
        # Check if the name has an asterisk (project leader)
        has_asterisk = '*' in name_part
        project_leader = None
        
        if has_asterisk:
            # Extract the name with PhD but without the asterisk
            project_leader = name_part.replace('*', '').strip()
            # Clean up any extra spaces
            project_leader = re.sub(r'\s+', ' ', project_leader)
            # Remove trailing comma if present
            project_leader = re.sub(r',\s*$', '', project_leader)
        else:
            # If no asterisk, the first author is the project leader by default
            # Keep the name as is (with PhD if present)
            project_leader = name_part.strip()
        
        # Clean the name for the authors list (remove PhD and asterisk for clean list)
        clean_name = name_part
        # Remove PhD, Dr., etc. for clean display
        clean_name = re.sub(r',\s*(?:PhD|Dr\.|Ph\.D\.)\s*', '', clean_name, flags=re.IGNORECASE)
        clean_name = re.sub(r'\s*(?:PhD|Dr\.|Ph\.D\.)\s*,?\s*', '', clean_name, flags=re.IGNORECASE)
        # Remove asterisk
        clean_name = clean_name.replace('*', '').strip()
        
        # Parse authors list
        authors_list = []
        if clean_name:
            authors_list.append(clean_name)
        
        # If there are multiple authors (separated by ;), split them
        if ';' in clean_name:
            parts = clean_name.split(';')
            authors_list = []
            for part in parts:
                part = part.strip()
                if part:
                    authors_list.append(part)
        
        # Build the full text with affiliation
        if affiliation_part:
            full_text = f"{name_part} / {affiliation_part}"
        else:
            full_text = authors_text
        
        return {
            'full_text': full_text,
            'list': authors_list,
            'project_leader': project_leader,
            'affiliation': affiliation_part if affiliation_part else None
        }
    
    def extract_corresponding_author(self):
        """Extract corresponding author - robustly handles typos/space breaks within emails."""
        if not self.text:
            return None
        
        # Normalize text: convert all newlines and multiple spaces to single spaces
        normalized_text = re.sub(r'\s+', ' ', self.text)
        
        # Look for the section label and split by "|" to isolate the value part
        pattern = r'3\.\s*Name\s+and\s+Email\s+Address\s+of\s+Corresponding\s+Author\s*\|\s*([^|]+)'
        match = re.search(pattern, normalized_text, re.IGNORECASE)
        
        if match:
            section_text = match.group(1).strip()
            
            # Regex to find the email address (ALLOWS SPACES AROUND THE DOT, e.g. "gmail. com")
            email_match = re.search(r'([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\s*\.\s*[a-zA-Z]{2,})', section_text)
            if email_match:
                email = re.sub(r'\s+', '', email_match.group(1).strip()) # Remove ALL spaces to fix typo
                # Extract name (everything before the email, clean up commas)
                name = section_text[:email_match.start()].strip()
                name = re.sub(r'[,\s]+$', '', name)
                
                if name and email:
                    return {
                        'full': f"{name} {email}".strip(),
                        'name': name,
                        'email': email
                    }
        
        # Fallback to generic regex if not found in the table
        email_pattern = r'([A-Z][A-Za-z\s.,]+(?:PhD|Dr\.)?)\s+([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\s*\.\s*[a-zA-Z]{2,})'
        match = re.search(email_pattern, normalized_text, re.IGNORECASE)
        if match:
            name = match.group(1).strip()
            email = re.sub(r'\s+', '', match.group(2).strip())
            return {
                'full': f"{name} {email}".strip(),
                'name': name,
                'email': email
            }
        
        return None
    
    def _get_highlighted_paragraphs(self):
        """Use python-docx to find paragraphs with yellow highlight or shading (including inside TABLES)."""
        try:
            from docx import Document
            import tempfile
            import os

            with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp_file:
                tmp_file.write(self.file_buffer)
                tmp_path = tmp_file.name

            doc = Document(tmp_path)
            highlighted_paragraphs = []

            # Helper function to check a single paragraph
            def check_paragraph(paragraph):
                for run in paragraph.runs:
                    # Check for text highlight
                    if run.font.highlight_color is not None:
                        if str(run.font.highlight_color) == 'YELLOW' or 'yellow' in str(run.font.highlight_color).lower():
                            highlighted_paragraphs.append(paragraph.text.strip())
                            return True
                    
                    # Check for shading (background color) in the run's XML
                    rPr = run._r.rPr
                    if rPr is not None:
                        shd = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
                        if shd is not None:
                            fill = shd.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill')
                            if fill and fill.lower() in ['ffff00', 'yellow']:
                                highlighted_paragraphs.append(paragraph.text.strip())
                                return True
                
                # Check paragraph shading
                pPr = paragraph._p.pPr
                if pPr is not None:
                    shd = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
                    if shd is not None:
                        fill = shd.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill')
                        if fill and fill.lower() in ['ffff00', 'yellow']:
                            highlighted_paragraphs.append(paragraph.text.strip())
                            return True
                            
                return False

            # 1. Check regular paragraphs
            for paragraph in doc.paragraphs:
                check_paragraph(paragraph)

            # 2. Check paragraphs inside TABLES (this is where your template is!)
            def recurse_tables(tables):
                for table in tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                check_paragraph(paragraph)
                            # Recursively check nested tables
                            if cell.tables:
                                recurse_tables(cell.tables)

            recurse_tables(doc.tables)

            os.unlink(tmp_path)
            return highlighted_paragraphs
        except Exception as e:
            print(f"Warning: Could not detect highlights in DOCX: {e}")
            return []

    def extract_paper_category(self):
        """Extract paper category - Handles [X], [✓], [/], and other markups, then fallback to highlighted color."""
        if not self.text:
            return None
        
        categories = ["Completed", "Ongoing"]
        
        # 1. Existing text checks. Remove newlines from text first to avoid line break issues.
        normalized_text = re.sub(r'\s+', ' ', self.text)
        for cat in categories:
            if re.search(r'\[(x|X|✓|√|✔|v|V|/)\]\s*' + cat, normalized_text, re.IGNORECASE):
                return f"{cat} Extension Project Paper"

        # 2. Fallback: Check highlighted paragraphs (only if text check failed)
        highlighted_paragraphs = self._get_highlighted_paragraphs()
        for text in highlighted_paragraphs:
            for cat in categories:
                if cat.lower() in text.lower():
                    return f"{cat} Extension Project Paper"

        return None

    def extract_thematic_area(self):
        """Extract thematic area - Handles [X], [✓], [/], and other markups, then fallback to highlighted color."""
        if not self.text:
            return None

        thematic_areas = [
            'Food Production, Agriculture, Fisheries, and Natural Resource Systems',
            'Health, Nutrition, Wellness, and Community Care',
            'Education, Literacy, Skills Development, and Lifelong Learning',
            'Livelihood, Entrepreneurship, Cooperatives, MSMEs, and Local Economic Development',
            'Environment, Climate Action, Disaster Risk Reduction, and Community Resilience'
        ]

        # 1. Existing text checks. Remove newlines from text first to avoid line break issues.
        normalized_text = re.sub(r'\s+', ' ', self.text)
        for area in thematic_areas:
            if re.search(r'\[(x|X|✓|√|✔|v|V|/)\]\s*' + re.escape(area), normalized_text, re.IGNORECASE):
                return area

        # 2. Fallback: Check highlighted paragraphs (only if text check failed)
        highlighted_paragraphs = self._get_highlighted_paragraphs()
        for text in highlighted_paragraphs:
            text_clean = re.sub(r'^[\[\]xX✓√✔vV/\s]+', '', text).strip()
            for area in thematic_areas:
                if area.lower() in text_clean.lower():
                    return area

        return None
    
    def extract_theme(self):
        """Extract theme - get the value after Theme:"""
        if not self.text:
            return None
        
        theme_match = re.search(r'Theme[:\s]*([^\n]+(?:\s*[^\n]+)*?)(?=\s*(?:Instructions|A\.|1\.|Paper Information|$))', self.text, re.IGNORECASE | re.DOTALL)
        if theme_match:
            theme = theme_match.group(1).strip()
            theme = re.sub(r'\s*(?:Instructions|A\.|1\.|Paper Information).*$', '', theme, flags=re.IGNORECASE)
            theme = re.sub(r'^Theme[:\s]*', '', theme, flags=re.IGNORECASE)
            if theme and len(theme) > 10:
                return theme
        
        return None
    
    def extract_all(self):
        """Extract all fields"""
        try:
            if not self.text:
                self.extract_text()
            
            if not self.text:
                return None
            
            extracted_data = {
                'title': self.extract_title(),
                'title_english': self.extract_title_english(),
                'authors_data': self.extract_authors(),
                'corresponding_author': self.extract_corresponding_author(),
                'paper_category': self.extract_paper_category(),
                'thematic_area': self.extract_thematic_area(),
                'theme': self.extract_theme()
            }
            
            extracted_data = {k: v for k, v in extracted_data.items() if v is not None}
            
            return extracted_data
            
        except Exception as e:
            print(f"Error extracting data: {e}")
            import traceback
            traceback.print_exc()
            return None

    @staticmethod
    def extract_from_file_buffer(file_buffer, filename=None):
        """Static method to extract data from file buffer (PDF or DOCX)"""
        extractor = DOCSExtractor(file_buffer, filename)
        return extractor.extract_all()